"""EXP-1 … EXP-5: the three Excel exports download formatted workbooks that respect the list filters."""

import io
from datetime import UTC, datetime
from decimal import Decimal
from typing import Any

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet
from sqlalchemy.orm import Session

from app.modules.exports import service as exports_service
from app.modules.exports.xlsx import XLSX_MEDIA_TYPE
from app.modules.invoices import service as invoices_service
from app.modules.stock.models import MovementType
from app.modules.users.models import Profile
from tests.factories import create_category, create_invoice, create_movement, create_product


def download(client: TestClient, headers: dict[str, str], path: str, **params: Any) -> tuple[str, Worksheet]:
    """Calls an export endpoint and returns (filename, the workbook's only sheet)."""
    response = client.get(f"/api/v1/exports/{path}", params=params, headers=headers)
    assert response.status_code == 200, response.text
    assert response.headers["content-type"] == XLSX_MEDIA_TYPE
    disposition = response.headers["content-disposition"]
    assert disposition.startswith("attachment; filename=")
    workbook = load_workbook(io.BytesIO(response.content))
    assert len(workbook.worksheets) == 1
    return disposition.split('filename="')[1].rstrip('"'), workbook.worksheets[0]


def values(sheet: Worksheet) -> list[list[Any]]:
    """Every row's cell values, header included."""
    return [list(row) for row in sheet.iter_rows(values_only=True)]


# ── Products (EXP-1) ────────────────────────────────────────────────────────


def test_products_export_is_formatted(client: TestClient, db: Session, staff_headers: dict[str, str]) -> None:
    tools = create_category(db, name="Tools")
    power = create_category(db, name="Power Tools", parent=tools)
    create_product(
        db, category=power, name="Cordless drill", sku="TL-D18", cost_price="1000", sell_price="1449.50",
        quantity=4, min_quantity=5,
    )  # fmt: skip
    create_product(db, category=tools, name="Hammer", sku="TL-HAM", cost_price="20", quantity=30)

    filename, sheet = download(client, staff_headers, "products")

    assert filename.startswith("products_") and filename.endswith(".xlsx")
    assert values(sheet) == [
        ["Name", "SKU", "Category", "Cost price", "Sell price", "Quantity", "Minimum",
         "Stock value (cost)", "Low stock", "Active"],
        ["Cordless drill", "TL-D18", "Tools › Power Tools", 1000, 1449.5, 4, 5, 4000, "Yes", "Yes"],
        ["Hammer", "TL-HAM", "Tools", 20, 15, 30, 0, 600, "No", "Yes"],
    ]  # fmt: skip
    # EXP-4: bold frozen header, currency format, autofilter, sized columns.
    assert sheet["A1"].font.b
    assert sheet.freeze_panes == "A2"
    assert sheet.auto_filter.ref == "A1:J3"
    assert '"SAR" #,##0.00' in sheet["D2"].number_format
    assert sheet.column_dimensions["C"].width >= len("Tools › Power Tools")


def test_products_export_respects_filters(
    client: TestClient, db: Session, admin_headers: dict[str, str]
) -> None:
    """EXP-5: the same search / category / low-stock / sort filters as the list."""
    paint = create_category(db, name="Paint")
    create_product(db, category=paint, name="White paint", sku="PT-W", quantity=2, min_quantity=5)
    create_product(db, category=paint, name="Black paint", sku="PT-B", quantity=50, min_quantity=5)
    create_product(db, name="Cement", sku="BM-C", quantity=1, min_quantity=5)
    create_product(db, category=paint, name="Old paint", sku="PT-O", is_active=False)

    _, sheet = download(
        client, admin_headers, "products", category_id=paint.id, sort="quantity", order="desc"
    )
    assert [row[0] for row in values(sheet)[1:]] == ["Black paint", "White paint"]

    _, sheet = download(client, admin_headers, "products", low_stock=True, search="paint")
    assert [row[0] for row in values(sheet)[1:]] == ["White paint"]

    _, sheet = download(client, admin_headers, "products", include_inactive=True, search="pt-")
    assert {row[0] for row in values(sheet)[1:]} == {"White paint", "Black paint", "Old paint"}


def test_text_that_looks_like_a_formula_stays_text(
    client: TestClient, db: Session, admin_headers: dict[str, str]
) -> None:
    """A user-typed name starting with "=" must not become a live Excel formula."""
    create_product(db, name='=HYPERLINK("http://evil.test","Click")', sku="BAD-1")
    _, sheet = download(client, admin_headers, "products")
    cell = sheet["A2"]
    assert cell.value == '=HYPERLINK("http://evil.test","Click")'
    assert cell.data_type == "s"


def test_empty_export_has_just_the_header(client: TestClient, admin_headers: dict[str, str]) -> None:
    _, sheet = download(client, admin_headers, "products", search="nothing-matches")
    assert len(values(sheet)) == 1


def test_export_too_large_for_one_sheet(
    client: TestClient, db: Session, admin_headers: dict[str, str], monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setattr(exports_service, "MAX_DATA_ROWS", 1)
    create_product(db)
    create_product(db)
    response = client.get("/api/v1/exports/products", headers=admin_headers)
    assert response.status_code == 422
    assert response.json()["code"] == "EXPORT_TOO_LARGE"


@pytest.mark.parametrize("path", ["products", "movements", "invoices"])
def test_exports_need_a_signed_in_user(client: TestClient, path: str) -> None:
    response = client.get(f"/api/v1/exports/{path}")
    assert response.status_code == 401
    assert response.json()["code"] == "NOT_AUTHENTICATED"


# ── Movements (EXP-2) ───────────────────────────────────────────────────────


def test_movements_export_for_a_date_range(
    client: TestClient, db: Session, staff: Profile, staff_headers: dict[str, str]
) -> None:
    product = create_product(db, name="Cement", sku="BM-C")
    # 22:30 UTC on 30 Sep is 01:30 on 1 Oct in Riyadh (UTC+3): inside the range.
    create_movement(
        db, product, staff, quantity=10, note="Delivery", created_at=datetime(2026, 9, 30, 22, 30, tzinfo=UTC)
    )
    create_movement(
        db, product, staff, type=MovementType.OUT, quantity=-3, note="=1+1 damaged",
        created_at=datetime(2026, 10, 5, 9, 0, tzinfo=UTC),
    )  # fmt: skip
    # Outside the range.
    create_movement(db, product, staff, quantity=1, created_at=datetime(2026, 9, 30, 20, 0, tzinfo=UTC))

    filename, sheet = download(
        client,
        staff_headers,
        "movements",
        start_date="2026-10-01",
        end_date="2026-10-31",
        page=2,
        page_size=1,
    )

    assert filename == "movements_2026-10-01_2026-10-31.xlsx"
    # Newest first; dates in business-local time; paging parameters are ignored.
    assert values(sheet) == [
        ["Date", "Product", "SKU", "Type", "Quantity", "Balance", "Note", "Invoice", "By"],
        [datetime(2026, 10, 5, 12, 0), "Cement", "BM-C", "OUT", -3, 7, "=1+1 damaged", None, "Staff"],
        [datetime(2026, 10, 1, 1, 30), "Cement", "BM-C", "IN", 10, 10, "Delivery", None, "Staff"],
    ]
    assert sheet["A2"].number_format == "yyyy-mm-dd hh:mm"
    assert sheet["G2"].data_type == "s"

    _, sheet = download(client, staff_headers, "movements", type="IN", start_date="2026-10-01")
    assert [row[3] for row in values(sheet)[1:]] == ["IN"]


def test_movements_export_shows_the_invoice(
    client: TestClient, db: Session, staff: Profile, staff_headers: dict[str, str]
) -> None:
    product = create_product(db, quantity=5)
    invoice = create_invoice(db, staff, [(product, 2)])
    filename, sheet = download(client, staff_headers, "movements", type="SALE")
    assert filename.startswith("movements_all_")
    assert values(sheet)[1][3:8] == ["SALE", -2, 3, None, invoice.invoice_number]


# ── Invoices (EXP-3) ────────────────────────────────────────────────────────


def test_invoices_export_respects_filters(
    client: TestClient, db: Session, admin: Profile, staff: Profile, admin_headers: dict[str, str]
) -> None:
    product = create_product(db, sell_price="100", quantity=50)
    other = create_product(db, sell_price="10", quantity=50)
    walk_in = create_invoice(
        db,
        staff,
        [(product, 2), (other, 1)],
        discount="10",
        created_at=datetime(2026, 10, 2, 9, 0, tzinfo=UTC),
    )
    named = create_invoice(
        db, staff, [(product, 1)], customer_name="Saad Contracting",
        created_at=datetime(2026, 10, 3, 9, 0, tzinfo=UTC),
    )  # fmt: skip
    invoices_service.cancel_invoice(db, admin, named.id)
    create_invoice(db, staff, [(product, 1)], created_at=datetime(2026, 9, 1, 9, 0, tzinfo=UTC))

    filename, sheet = download(
        client, admin_headers, "invoices", start_date="2026-10-01", end_date="2026-10-31"
    )

    assert filename == "invoices_2026-10-01_2026-10-31.xlsx"
    assert values(sheet) == [
        ["Invoice", "Date", "Customer", "Items", "Subtotal", "Discount", "Tax", "Total", "Status", "Created by"],
        [named.invoice_number, datetime(2026, 10, 3, 12, 0), "Saad Contracting", 1, 100, 0, 0, 100, "CANCELLED", "Staff"],
        [walk_in.invoice_number, datetime(2026, 10, 2, 12, 0), "Walk-in customer", 2, 210, 10, 0, 200, "COMPLETED", "Staff"],
    ]  # fmt: skip

    # No date range: both completed walk-in sales (October and September), newest first.
    _, sheet = download(client, admin_headers, "invoices", status="COMPLETED", search="walk-in")
    rows = values(sheet)[1:]
    assert [(row[2], row[8]) for row in rows] == [("Walk-in customer", "COMPLETED")] * 2
    assert rows[0][0] == walk_in.invoice_number


def test_money_keeps_its_cents(
    client: TestClient, db: Session, staff: Profile, admin_headers: dict[str, str]
) -> None:
    product = create_product(db, sell_price="19.99", quantity=5)
    create_invoice(db, staff, [(product, 3)])
    _, sheet = download(client, admin_headers, "invoices")
    assert Decimal(str(sheet["H2"].value)) == Decimal("59.97")
