"""Workbook formatting rules (EXP-4), checked without a database."""

import io
from datetime import UTC, datetime
from decimal import Decimal
from zoneinfo import ZoneInfo

from openpyxl import load_workbook

from app.modules.exports.xlsx import Column, Kind, build_workbook

RIYADH = ZoneInfo("Asia/Riyadh")


def test_values_are_formatted_by_kind() -> None:
    columns = [
        Column("Name"),
        Column("Price", Kind.MONEY),
        Column("Change", Kind.SIGNED),
        Column("At", Kind.DATETIME),
        Column("Active", Kind.FLAG),
    ]
    rows = [("Drill", Decimal("1449.50"), -3, datetime(2026, 10, 19, 21, 30, tzinfo=UTC), True)]
    content = build_workbook("Things", columns, rows, currency="AED", tz=RIYADH)

    sheet = load_workbook(io.BytesIO(content))["Things"]
    assert [c.value for c in sheet[2]] == ["Drill", 1449.5, -3, datetime(2026, 10, 20, 0, 30), "Yes"]
    assert sheet["B2"].number_format.startswith('"AED" #,##0.00')
    assert sheet["C2"].number_format == "+#,##0;-#,##0;0"
    assert sheet["D2"].number_format == "yyyy-mm-dd hh:mm"
    assert all(cell.font.b for cell in sheet[1])


def test_missing_values_stay_empty() -> None:
    content = build_workbook(
        "S", [Column("Note"), Column("Total", Kind.MONEY)], [(None, None)], currency="SAR", tz=UTC
    )
    sheet = load_workbook(io.BytesIO(content))["S"]
    assert [c.value for c in sheet[2]] == [None, None]


def test_column_widths_fit_the_content_within_limits() -> None:
    columns = [Column("A"), Column("Long")]
    rows = [("x" * 20, "y" * 500)]
    sheet = load_workbook(io.BytesIO(build_workbook("S", columns, rows, currency="SAR", tz=UTC)))["S"]
    assert sheet.column_dimensions["A"].width == 22
    assert sheet.column_dimensions["B"].width == 60
